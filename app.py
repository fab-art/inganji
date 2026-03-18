from __future__ import annotations

import csv
import io
from dataclasses import asdict, dataclass
from difflib import SequenceMatcher
from datetime import datetime
from pathlib import Path
from typing import BinaryIO, Dict, Iterable, List, Optional

REQUIRED_FIELDS = ("patient_id", "patient_name", "visit_date")
KEYWORDS = {
    "patient_id": ["rama", "rama number", "patient_rama_number", "id"],
    "patient_name": ["name", "beneficiary", "patient_name"],
    "visit_date": ["date", "dispensing_date", "prescription_date", "visit_date"],
}


@dataclass(frozen=True)
class AnalysisConfig:
    max_file_size_mb: int = 100
    max_date_gap_days: int = 30
    fuzzy_name_threshold: float = 85.0


@dataclass(frozen=True)
class FileMetadata:
    file_name: str
    file_size_bytes: int
    row_count: int
    column_count: int
    min_date: Optional[datetime]
    max_date: Optional[datetime]


@dataclass(frozen=True)
class MappingSuggestion:
    field: str
    column: str
    score: float


@dataclass(frozen=True)
class MatchResult:
    pharmacy_index: int
    clinic_index: Optional[int]
    rama_match: bool
    name_similarity: float
    date_valid: bool
    confidence: float


@dataclass(frozen=True)
class Anomaly:
    pharmacy_index: int
    kind: str
    severity: str
    details: str


@dataclass(frozen=True)
class AnalysisResult:
    matches: List[MatchResult]
    anomalies: List[Anomaly]
    summary_rows: List[dict]
    stats: Dict[str, int]


class FileValidationError(ValueError):
    """Raised when file validation fails."""


class FileHandler:
    SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}

    def __init__(self, config: Optional[AnalysisConfig] = None) -> None:
        self.config = config or AnalysisConfig()

    def read_uploaded_file(self, filename: str, payload: bytes) -> List[dict]:
        suffix = Path(filename).suffix.lower()
        self._validate_extension(suffix)
        self._validate_file_size(len(payload))
        return self.read_binary(suffix, io.BytesIO(payload))

    def read_binary(self, suffix: str, handle: BinaryIO) -> List[dict]:
        if suffix == ".csv":
            text = handle.read().decode("utf-8-sig", errors="replace")
            return list(csv.DictReader(io.StringIO(text)))

        from openpyxl import load_workbook

        workbook = load_workbook(handle, read_only=True, data_only=True)
        worksheet = workbook.active
        rows = list(worksheet.iter_rows(values_only=True))
        if not rows:
            return []

        headers = [str(header) if header is not None else "" for header in rows[0]]
        return [
            {headers[index]: row[index] for index in range(len(headers))}
            for row in rows[1:]
        ]

    def build_metadata(
        self,
        file_name: str,
        rows: List[dict],
        file_size_bytes: int,
        date_columns: Optional[Iterable[str]] = None,
    ) -> FileMetadata:
        columns = list(rows[0].keys()) if rows else []
        candidates = list(date_columns or []) or [column for column in columns if "date" in column.lower()]

        min_date = None
        max_date = None
        for row in rows:
            for column in candidates:
                parsed = self.parse_date(row.get(column))
                if not parsed:
                    continue
                min_date = parsed if min_date is None else min(min_date, parsed)
                max_date = parsed if max_date is None else max(max_date, parsed)

        return FileMetadata(
            file_name=file_name,
            file_size_bytes=file_size_bytes,
            row_count=len(rows),
            column_count=len(columns),
            min_date=min_date,
            max_date=max_date,
        )

    def _validate_extension(self, suffix: str) -> None:
        if suffix not in self.SUPPORTED_EXTENSIONS:
            raise FileValidationError(f"Unsupported file type: {suffix}")

    def _validate_file_size(self, file_size_bytes: int) -> None:
        max_bytes = self.config.max_file_size_mb * 1024 * 1024
        if file_size_bytes > max_bytes:
            raise FileValidationError("File exceeds configured max size")

    @staticmethod
    def parse_date(value: object) -> Optional[datetime]:
        if value is None or value == "":
            return None
        if isinstance(value, datetime):
            return value

        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        return None


class ColumnMapper:
    def suggest_mapping(self, column_names: Iterable[str]) -> Dict[str, MappingSuggestion]:
        normalized_columns = [str(column).strip() for column in column_names]
        suggestions: Dict[str, MappingSuggestion] = {}

        for field, keywords in KEYWORDS.items():
            best_column = ""
            best_score = -1.0
            for column in normalized_columns:
                score = max(self._score_similarity(column.lower(), keyword.lower()) for keyword in keywords)
                if score > best_score:
                    best_column = column
                    best_score = score
            if best_column:
                suggestions[field] = MappingSuggestion(field=field, column=best_column, score=best_score)

        return suggestions

    def auto_detect_columns(self, column_names: Iterable[str]) -> Dict[str, str]:
        return {
            field: suggestion.column
            for field, suggestion in self.suggest_mapping(column_names).items()
            if suggestion.score >= 70.0
        }

    def validate_mapping(self, mapping: Dict[str, str]) -> bool:
        return all(mapping.get(field) for field in REQUIRED_FIELDS)

    @staticmethod
    def _score_similarity(left: str, right: str) -> float:
        return SequenceMatcher(None, left, right).ratio() * 100.0


class PatientMatcher:
    def __init__(self, config: Optional[AnalysisConfig] = None) -> None:
        self.config = config or AnalysisConfig()

    def match_by_rama(
        self,
        pharmacy_rows: List[dict],
        clinic_rows: List[dict],
        pharmacy_map: Dict[str, str],
        clinic_map: Dict[str, str],
    ) -> List[MatchResult]:
        clinic_index: Dict[str, List[tuple[int, dict]]] = {}
        for clinic_position, row in enumerate(clinic_rows):
            rama_value = self._normalize_identifier(row.get(clinic_map["patient_id"]))
            clinic_index.setdefault(rama_value, []).append((clinic_position, row))

        results: List[MatchResult] = []
        for pharmacy_position, row in enumerate(pharmacy_rows):
            rama_value = self._normalize_identifier(row.get(pharmacy_map["patient_id"]))
            if not rama_value or rama_value not in clinic_index:
                results.append(MatchResult(pharmacy_position, None, False, 0.0, False, 0.0))
                continue

            best_match: Optional[MatchResult] = None
            for clinic_position, candidate in clinic_index[rama_value]:
                similarity = self.fuzzy_name_match(
                    str(row.get(pharmacy_map["patient_name"], "")),
                    str(candidate.get(clinic_map["patient_name"], "")),
                )
                pharmacy_date = FileHandler.parse_date(row.get(pharmacy_map["visit_date"]))
                clinic_date = FileHandler.parse_date(candidate.get(clinic_map["visit_date"]))
                date_valid = bool(pharmacy_date and clinic_date and self.validate_date_sequence(pharmacy_date, clinic_date))
                confidence = self.calculate_confidence_score(True, similarity, date_valid)
                evaluated = MatchResult(
                    pharmacy_index=pharmacy_position,
                    clinic_index=clinic_position,
                    rama_match=True,
                    name_similarity=similarity,
                    date_valid=date_valid,
                    confidence=confidence,
                )
                if best_match is None or evaluated.confidence > best_match.confidence:
                    best_match = evaluated

            results.append(best_match or MatchResult(pharmacy_position, None, False, 0.0, False, 0.0))
        return results

    def validate_date_sequence(self, dispensing_date: datetime, clinic_date: datetime) -> bool:
        delta_days = (dispensing_date - clinic_date).days
        return 0 <= delta_days <= self.config.max_date_gap_days

    def calculate_confidence_score(self, rama_match: bool, name_similarity: float, date_valid: bool) -> float:
        score = 70.0 if rama_match else 0.0
        score += min(name_similarity, 100.0) * 0.2
        score += 10.0 if date_valid else 0.0
        return round(min(score, 100.0), 2)

    @staticmethod
    def fuzzy_name_match(left: str, right: str) -> float:
        left_words = set(left.lower().split())
        right_words = set(right.lower().split())
        if not left_words and not right_words:
            return 100.0
        if not left_words or not right_words:
            return 0.0
        overlap = len(left_words & right_words)
        total = len(left_words | right_words)
        return 100.0 * overlap / total

    @staticmethod
    def _normalize_identifier(value: object) -> str:
        return str(value).strip().upper() if value is not None else ""


class AnomalyDetector:
    def find_unmatched_pharmacy_visits(self, matches: List[MatchResult]) -> List[Anomaly]:
        return [
            Anomaly(
                pharmacy_index=match.pharmacy_index,
                kind="unmatched_pharmacy_visit",
                severity="Critical",
                details="No clinic record found for RAMA number",
            )
            for match in matches
            if not match.rama_match or match.clinic_index is None
        ]

    def detect_date_anomalies(self, matches: List[MatchResult]) -> List[Anomaly]:
        return [
            Anomaly(
                pharmacy_index=match.pharmacy_index,
                kind="invalid_date_sequence",
                severity="High",
                details="Clinic date is after dispensing or exceeds configured date window",
            )
            for match in matches
            if match.clinic_index is not None and not match.date_valid
        ]

    def check_name_discrepancies(self, matches: List[MatchResult], threshold: float) -> List[Anomaly]:
        return [
            Anomaly(
                pharmacy_index=match.pharmacy_index,
                kind="name_mismatch",
                severity="Medium",
                details=f"Name similarity below threshold: {match.name_similarity:.2f}",
            )
            for match in matches
            if match.clinic_index is not None and match.name_similarity < threshold
        ]


class ReportGenerator:
    def generate_summary_report(self, anomalies: List[Anomaly]) -> List[dict]:
        return [asdict(anomaly) for anomaly in anomalies]

    def generate_statistics(self, anomalies: List[Anomaly], total_analyzed: int) -> Dict[str, int]:
        return {
            "total_analyzed": total_analyzed,
            "total_anomalies": len(anomalies),
            "severity_critical": sum(1 for item in anomalies if item.severity == "Critical"),
            "severity_high": sum(1 for item in anomalies if item.severity == "High"),
            "severity_medium": sum(1 for item in anomalies if item.severity == "Medium"),
            "severity_low": sum(1 for item in anomalies if item.severity == "Low"),
        }

    def export_csv(self, rows: List[dict]) -> bytes:
        if not rows:
            return b""
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)
        return buffer.getvalue().encode("utf-8")


class AnalysisService:
    def __init__(self, config: Optional[AnalysisConfig] = None) -> None:
        self.config = config or AnalysisConfig()
        self.matcher = PatientMatcher(self.config)
        self.detector = AnomalyDetector()
        self.reporter = ReportGenerator()

    def run(
        self,
        pharmacy_rows: List[dict],
        clinic_rows: List[dict],
        pharmacy_map: Dict[str, str],
        clinic_map: Dict[str, str],
    ) -> AnalysisResult:
        matches = self.matcher.match_by_rama(pharmacy_rows, clinic_rows, pharmacy_map, clinic_map)
        anomalies = [
            *self.detector.find_unmatched_pharmacy_visits(matches),
            *self.detector.detect_date_anomalies(matches),
            *self.detector.check_name_discrepancies(matches, self.config.fuzzy_name_threshold),
        ]
        summary_rows = self.reporter.generate_summary_report(anomalies)
        stats = self.reporter.generate_statistics(anomalies, total_analyzed=len(pharmacy_rows))
        return AnalysisResult(matches=matches, anomalies=anomalies, summary_rows=summary_rows, stats=stats)


def render_app() -> None:
    import streamlit as st

    st.set_page_config(page_title="Patient Cross-Checking System", layout="wide")
    st.title("Patient Cross-Checking System")
    st.caption("Single-file Streamlit deployment for pharmacy and clinic visit validation")

    file_handler = FileHandler()
    mapper = ColumnMapper()
    analyzer = AnalysisService()

    with st.sidebar:
        st.header("Inputs")
        pharmacy_file = st.file_uploader("Pharmacy file", type=["csv", "xlsx", "xls"])
        clinic_file = st.file_uploader("Clinic file", type=["csv", "xlsx", "xls"])

    if not pharmacy_file or not clinic_file:
        st.info("Upload both pharmacy and clinic files to begin analysis.")
        return

    try:
        pharmacy_payload = pharmacy_file.getvalue()
        clinic_payload = clinic_file.getvalue()
        pharmacy_rows = file_handler.read_uploaded_file(pharmacy_file.name, pharmacy_payload)
        clinic_rows = file_handler.read_uploaded_file(clinic_file.name, clinic_payload)
    except FileValidationError as exc:
        st.error(str(exc))
        return

    pharmacy_metadata = file_handler.build_metadata(pharmacy_file.name, pharmacy_rows, len(pharmacy_payload))
    clinic_metadata = file_handler.build_metadata(clinic_file.name, clinic_rows, len(clinic_payload))

    st.success("Files loaded successfully")
    overview_left, overview_right = st.columns(2)
    overview_left.metric("Pharmacy rows", pharmacy_metadata.row_count)
    overview_right.metric("Clinic rows", clinic_metadata.row_count)

    def mapping_ui(prefix: str, columns: List[str], detected: Dict[str, str]) -> Dict[str, str]:
        st.subheader(f"{prefix} column mapping")
        mapping: Dict[str, str] = {}
        for field in REQUIRED_FIELDS:
            options = [""] + columns
            detected_value = detected.get(field, "")
            index = options.index(detected_value) if detected_value in options else 0
            mapping[field] = st.selectbox(f"{prefix} {field}", options, index=index, key=f"{prefix}_{field}")
        return mapping

    pharmacy_columns = list(pharmacy_rows[0].keys()) if pharmacy_rows else []
    clinic_columns = list(clinic_rows[0].keys()) if clinic_rows else []
    pharmacy_map = mapping_ui("Pharmacy", pharmacy_columns, mapper.auto_detect_columns(pharmacy_columns))
    clinic_map = mapping_ui("Clinic", clinic_columns, mapper.auto_detect_columns(clinic_columns))

    if not st.button("Run analysis", type="primary"):
        return

    if not mapper.validate_mapping(pharmacy_map) or not mapper.validate_mapping(clinic_map):
        st.error("Please map patient_id, patient_name, and visit_date for both files.")
        return

    result = analyzer.run(pharmacy_rows, clinic_rows, pharmacy_map, clinic_map)

    stat_1, stat_2, stat_3, stat_4 = st.columns(4)
    stat_1.metric("Analyzed", result.stats["total_analyzed"])
    stat_2.metric("Anomalies", result.stats["total_anomalies"])
    stat_3.metric("Critical", result.stats["severity_critical"])
    stat_4.metric("High", result.stats["severity_high"])

    st.subheader("Anomalies")
    if not result.summary_rows:
        st.info("No anomalies found for the selected files and mapping.")
        return

    st.dataframe(result.summary_rows, use_container_width=True)
    st.download_button(
        "Download anomalies (CSV)",
        analyzer.reporter.export_csv(result.summary_rows),
        file_name="anomalies_report.csv",
        mime="text/csv",
    )


if __name__ == "__main__":
    render_app()
