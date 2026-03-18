from __future__ import annotations

import csv
import io
from datetime import datetime
from pathlib import Path
from typing import BinaryIO, Iterable, List, Optional

from .models import AnalysisConfig, FileMetadata


class FileValidationError(ValueError):
    """Raised when file validation fails."""


class FileHandler:
    """Handles file validation and ingestion for CSV/Excel sources."""

    SUPPORTED_EXTENSIONS = {".csv", ".xlsx", ".xls"}

    def __init__(self, config: Optional[AnalysisConfig] = None) -> None:
        self.config = config or AnalysisConfig()

    def validate_file(self, file_path: str | Path) -> bool:
        path = Path(file_path)
        if not path.exists() or not path.is_file():
            raise FileValidationError(f"File not found: {path}")
        self._validate_extension(path.suffix)

        max_bytes = self.config.max_file_size_mb * 1024 * 1024
        if path.stat().st_size > max_bytes:
            raise FileValidationError("File exceeds configured max size")
        return True

    def detect_encoding(self, file_path: str | Path) -> str:
        path = Path(file_path)
        if path.suffix.lower() != ".csv":
            return "binary"
        for enc in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                with path.open("r", encoding=enc) as f:
                    f.read(2048)
                return enc
            except UnicodeDecodeError:
                pass
        return "latin-1"

    def read_file(self, file_path: str | Path) -> List[dict]:
        self.validate_file(file_path)
        path = Path(file_path)
        if path.suffix.lower() == ".csv":
            encoding = self.detect_encoding(path)
            with path.open("r", encoding=encoding, newline="") as f:
                return list(csv.DictReader(f))

        with path.open("rb") as handle:
            return self.read_binary(path.suffix.lower(), handle)

    def read_uploaded_file(self, filename: str, payload: bytes) -> List[dict]:
        suffix = Path(filename).suffix.lower()
        self._validate_extension(suffix)

        max_bytes = self.config.max_file_size_mb * 1024 * 1024
        if len(payload) > max_bytes:
            raise FileValidationError("File exceeds configured max size")

        return self.read_binary(suffix, io.BytesIO(payload))

    def read_binary(self, suffix: str, handle: BinaryIO) -> List[dict]:
        if suffix == ".csv":
            text = handle.read().decode("utf-8-sig", errors="replace")
            return list(csv.DictReader(io.StringIO(text)))

        try:
            from openpyxl import load_workbook
        except Exception as exc:  # pragma: no cover - optional dependency
            raise FileValidationError("openpyxl is required to read Excel files") from exc

        wb = load_workbook(handle, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        return [{headers[i]: row[i] for i in range(len(headers))} for row in rows[1:]]

    def get_file_metadata(
        self,
        file_path: str | Path,
        date_columns: Optional[Iterable[str]] = None,
    ) -> FileMetadata:
        rows = self.read_file(file_path)
        path = Path(file_path)
        return self.build_metadata(path, rows, path.stat().st_size, date_columns)

    def build_metadata(
        self,
        file_path: Path,
        rows: List[dict],
        file_size_bytes: int,
        date_columns: Optional[Iterable[str]] = None,
    ) -> FileMetadata:
        columns = list(rows[0].keys()) if rows else []
        candidate_columns = list(date_columns or []) or [c for c in columns if "date" in c.lower()]

        min_date = None
        max_date = None
        for row in rows:
            for col in candidate_columns:
                parsed = self._parse_date(row.get(col))
                if not parsed:
                    continue
                min_date = parsed if min_date is None else min(min_date, parsed)
                max_date = parsed if max_date is None else max(max_date, parsed)

        return FileMetadata(
            file_path=file_path,
            file_size_bytes=file_size_bytes,
            row_count=len(rows),
            column_count=len(columns),
            min_date=min_date,
            max_date=max_date,
        )

    def _validate_extension(self, suffix: str) -> None:
        if suffix.lower() not in self.SUPPORTED_EXTENSIONS:
            raise FileValidationError(f"Unsupported file type: {suffix}")

    @staticmethod
    def _parse_date(value: object) -> Optional[datetime]:
        if value is None or value == "":
            return None
        if isinstance(value, datetime):
            return value
        text = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        return None
